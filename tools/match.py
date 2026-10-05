#!/usr/bin/env python3
"""Pair founders from their rankings. Reads files, writes a spreadsheet; sends nothing anywhere.

    python tools/match.py private/responses.csv

responses.csv is the Google Sheet export (Timestamp, Response). Each Response is JSON {"t": token, "c": [ids]}.
Needs private/roster.csv from make_links.py. Writes private/matching_results.xlsx with three sheets:
recommended pairings, every possible pair with its score, and notes.

Method: a pair's score is rankA(B)^2 + rankB(A)^2 (lower is better, symmetric). Lists can be any length;
someone left off a list counts as rank N (the group size), worse than any ranked position.
Mutual preference comes first: among pairs where both people ranked each other, the best-scoring pair is
locked, those two are removed, and this repeats (ties at the same score are resolved so that as many of the
tied pairs as possible are kept). A mutual pair is never split to improve the group total.
Whoever remains has no mutual option left; they are paired so the sum of their scores is as small as possible
(minimum-weight matching). With an odd count, the person left over joins the pair where they add the least.
"""
import argparse, csv, json
import networkx as nx
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("responses")
ap.add_argument("--out", default="private/matching_results.xlsx")
args = ap.parse_args()

roster = list(csv.DictReader(open("private/roster.csv")))
by_token = {r["token"]: int(r["id"]) for r in roster}
name = {int(r["id"]): r["name"] for r in roster if int(r["id"]) > 0}  # id 0 is the hidden test user
ids = sorted(name)

# latest valid submission per person; rows are in time order
prefs, bad = {}, 0
for row in csv.reader(open(args.responses)):
    try:
        d = json.loads(row[-1]); me = by_token[d["t"]]
    except (ValueError, KeyError, IndexError, TypeError):
        bad += 1; continue
    if me not in name:
        continue
    seen, clean = set(), []
    for c in d.get("c", []):
        if isinstance(c, int) and c in name and c != me and c not in seen:
            seen.add(c); clean.append(c)
    prefs[me] = clean

UNRANKED = len(ids)
rank = lambda a, b: prefs.get(a, []).index(b) + 1 if b in prefs.get(a, []) else UNRANKED
cost = lambda a, b: rank(a, b) ** 2 + rank(b, a) ** 2
label = lambda r: "not ranked" if r == UNRANKED else r

mutual = lambda a, b: rank(a, b) < UNRANKED and rank(b, a) < UNRANKED
levels = {}
for i, a in enumerate(ids):
    for b in ids[i + 1:]:
        if mutual(a, b):
            levels.setdefault(cost(a, b), []).append((a, b))
locked, taken = [], set()
for score in sorted(levels):
    T = nx.Graph((a, b) for a, b in levels[score] if a not in taken and b not in taken)
    for a, b in sorted(tuple(sorted(e)) for e in nx.max_weight_matching(T, maxcardinality=True)):
        locked.append((a, b)); taken.update((a, b))
rest = [p for p in ids if p not in taken]

G = nx.Graph()
G.add_weighted_edges_from((a, b, cost(a, b)) for i, a in enumerate(rest) for b in rest[i + 1:])
groups = [list(p) for p in locked] + [sorted(p) for p in nx.min_weight_matching(G)]

left = [p for p in ids if p not in {q for g in groups for q in g}]
for p in left:  # at most one, when the count is odd
    min((g for g in groups if len(g) == 2), key=lambda g: sum(cost(p, q) for q in g)).append(p)

partner = {p: [q for q in g if q != p] for g in groups for p in g}
worst = lambda p: max(rank(p, q) for q in partner[p])
blocking = [(a, b) for i, a in enumerate(ids) for b in ids[i + 1:]
            if b not in partner[a] and rank(a, b) < worst(a) and rank(b, a) < worst(b)]
group_cost = lambda g: sum(cost(a, b) for i, a in enumerate(g) for b in g[i + 1:])
groups.sort(key=group_cost)
together = {frozenset((a, b)) for g in groups for i, a in enumerate(g) for b in g[i + 1:]}
missing = [name[p] for p in ids if p not in prefs]

wb = Workbook()
def sheet(ws, title, header, rows, widths):
    ws.title = title
    ws.append(header)
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows: ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

rec = []
for n, g in enumerate(groups, 1):
    for i, a in enumerate(g):
        for b in g[i + 1:]:
            rec.append([n, name[a], name[b], label(rank(a, b)), label(rank(b, a)), cost(a, b),
                        "mutual" if mutual(a, b) else "not mutual"])
sheet(wb.active, "Recommended pairings",
      ["Group", "Person A", "Person B", "A ranked B", "B ranked A", "Score (lower is better)", "Note"], rec, [8, 26, 26, 12, 12, 22, 22])

allp = sorted(((cost(a, b), a, b) for i, a in enumerate(ids) for b in ids[i + 1:]), key=lambda t: (t[0], name[t[1]], name[t[2]]))
sheet(wb.create_sheet(), "All pair scores",
      ["Person A", "Person B", "A ranked B", "B ranked A", "Score (lower is better)", "Recommended"],
      [[name[a], name[b], label(rank(a, b)), label(rank(b, a)), c, "yes" if frozenset((a, b)) in together else ""] for c, a, b in allp],
      [26, 26, 12, 12, 22, 14])

notes = [["People", len(ids)], ["Submitted", len(prefs)], ["Mutual pairs (both ranked each other)", len(locked)],
         ["Total score of recommended pairings", sum(group_cost(g) for g in groups)],
         ["Did not submit", ", ".join(missing) or "none"],
         ["Pairs who would both prefer each other over their assigned partner (review by hand)",
          "; ".join(f"{name[a]} & {name[b]}" for a, b in blocking) or "none"],
         ["Score", "A's rank of B squared plus B's rank of A squared. 'Not ranked' counts as rank %d." % UNRANKED],
         ["Method", "Pairs where both people ranked each other are locked best score first; a mutual pair is never split to help the group total. Anyone left over is paired to give the lowest total score among them."]]
sheet(wb.create_sheet(), "Notes", ["Item", "Value"], notes, [70, 90])
wb.save(args.out)

print(f"wrote {args.out}: {len(groups)} groups, {len(allp)} possible pairs scored")
print(f"{len(ids)} people, {len(prefs)} submitted, {len(locked)} mutual pairs locked, {len(blocking)} pairs flagged for review, {len(missing)} did not submit")
